from celery import shared_task

#     # # Define the media directory
    #     # media_dir = os.path.join("/vol/web/media", 'xlsx')
    #     # os.makedirs(media_dir, exist_ok=True)  # Ensure the directory exists\
    #
    #     media_dir = os.path.join(settings.MEDIA_ROOT, 'xlsx')
    #     os.makedirs(media_dir, exist_ok=True)
    #
    #
    #     # Save the exported file to the mounted volume
    #     file_path = os.path.join(media_dir, file_name)  # Path to the mounted volume
    #     with default_storage.open(file_path, "wb") as file_content_file:
    #         file_content_file.write(file_content)
    #
    #     # Return the URL of the exported file
    #     file_url = os.path.join(settings.MEDIA_URL,'xlsx', file_name)  # Assuming MEDIA_URL is /media/
    #     return {
    #         'file_url': file_url,
    #         'file_name': file_name,
    #     }
    #
    # except Exception as e:
    #     return {'error': str(e)}




from hydrocarbures.celery import app
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage


@app.task
def exportLargeDataSet(export_format, queryset_data, request=None):
    """Legacy exporter used by some views/templates. Keeps compatibility."""
    try:
        # Lazy imports to avoid hard dependency at worker startup
        try:
            from django_tables2 import RequestConfig  # type: ignore
            from django_tables2.export.export import TableExport  # type: ignore
            from ads.tables import RapportBrut  # type: ignore
        except Exception as imp_err:
            return {'error': f"django-tables2 is required for exportLargeDataSet: {imp_err}"}

        table = RapportBrut(queryset_data)
        if request:
            RequestConfig(request, paginate={"per_page": 15}).configure(table)

        exporter = TableExport(export_format, table, exclude_columns=("actions"))
        file_content = exporter.export()

        import datetime as _dt
        current_datetime = _dt.datetime.now().strftime("%Y%m%d%H%M%S")
        file_name = f"table_{current_datetime}.{export_format}"
        file_path = f"xlsx/{file_name}"
        file_content_file = ContentFile(file_content)
        default_storage.save(file_path, file_content_file)
        full_file_url = default_storage.url(file_path)

        return {'file_url': full_file_url, 'file_name': file_name}
    except Exception as e:
        return {'error': str(e)}




@shared_task(bind=True)
def exportRapportBrutExcel(self, params: dict):
    """
    Build the queryset using provided filters, stream an Excel file, report progress,
    save to storage, and return {file_url, file_name}.

    Expected params keys:
      - date_from, date_to (YYYY-MM-DD or date-like)
      - entrepot_ville_id (int) — filters by frontiere_id (denormalized frontiere)
      - importateur (int), entrepot (int), produit (int)
      - immatriculation (str), declaration (str), numdos (str) — optional icontains filters
    """
    try:
        # Lazy imports to avoid heavy initialization at worker start
        from django.db.models import Q
        from django.utils import timezone
        from django.core.files.base import ContentFile
        from django.core.files.storage import default_storage
        from openpyxl import Workbook
        from enreg.models import Cargaison

        def g(key, default=None):
            return params.get(key, default)

        # -------------------------
        # Base queryset (denormalized fields)
        # -------------------------
        qs = Cargaison.objects.all().values(
            'idcargaison',
            'numdos',                # ✅ keep for Excel "#.DOSSIER"
            'numdossier',
            'numreq',
            'declaration',
            'nom_frontiere',
            'nom_entrepot',
            'nom_importateur',
            'immatriculation',
            'nom_produit',
            'dateheurecargaison',
            'requisitiondackdate',
            'date_echantillon',
            'date_reception_labo',
            'date_analyse',
            'date_inspection',
            'dateDechargement',
            'volume',
            'gov_total',
            'gsv_total',
            'mta_total',
            'mtv_total',
            'densite_inspection',
            'temperature_inspection',
        )

        # -------------------------
        # Filters
        # -------------------------
        adv = Q()

        date_from = g('date_from')
        date_to = g('date_to')
        if date_from and date_to:
            adv &= Q(dateheurecargaison__date__range=[date_from, date_to])
        elif date_from:
            adv &= Q(dateheurecargaison__date__gte=date_from)
        elif date_to:
            adv &= Q(dateheurecargaison__date__lte=date_to)

        v_id = g('entrepot_ville_id')
        if v_id is not None and str(v_id).strip() != '':
            try:
                # Use denormalized frontiere_id for performance
                adv &= Q(frontiere_id=int(v_id))
            except Exception:
                pass

        for key, field in [
            ('produit_id', 'produit_id'),
            ('importateur_id', 'importateur_id'),
            ('entrepot_id', 'entrepot_id'),
            ('produit', 'produit_id'),
            ('importateur', 'importateur_id'),
            ('entrepot', 'entrepot_id'),
        ]:
            val = g(key)
            try:
                if val is not None and str(val).strip() != '':
                    adv &= Q(**{field: int(val)})
            except Exception:
                pass

        # Optional text filters
        for key, field in [
            ('immatriculation', 'immatriculation__icontains'),
            ('declaration', 'declaration__icontains'),
        ]:
            val = g(key)
            if isinstance(val, str) and val.strip():
                adv &= Q(**{field: val.strip()})

        # Safe filter for dossier (search across multiple potential fields)
        numdos_val = g('numdos')
        if numdos_val:
            s_num = str(numdos_val).strip()
            if s_num:
                d_q = Q(numdossier__icontains=s_num) | Q(numreq__icontains=s_num)
                if s_num.isdigit():
                    d_q |= Q(numdos=int(s_num))
                adv &= d_q

        if adv:
            qs = qs.filter(adv)

        # -------------------------
        # Workbook
        # -------------------------
        wb = Workbook(write_only=True)
        ws = wb.create_sheet('Rapport')
        try:
            if len(wb.worksheets) > 1:
                wb.remove(wb.worksheets[0])
        except Exception:
            pass

        # ✅ removed duplicated "NUMDOS" column
        headers = [
            'DATE ENTREE','FRONTIERE','FOURNISSEUR','ENTREPOT','PRODUIT','VOL.DECL.',
            'IMMATR.','#.DECLARATION','#.DOSSIER','DATE REQUISITION','DATE ECHANTILLONNAGE',
            'DATE RECEPTION LABO','DATE D\'ANALYSE','DATE D\'INSPECTION','DATE DE DECHARGEMENT',
            'VOL JAUGE (GOV)','DENSITE @15','TEMPERATURE','VCF','MTA','MTV','GSV'
        ]
        ws.append(headers)

        total = qs.count()
        done = 0

        from entrepot.calculs import densite15, vcf
        from django.utils import timezone as _tz

        def _fmt_dt(val):
            try:
                import datetime as _dt
                if val is None:
                    return None
                if isinstance(val, _dt.datetime):
                    if _tz.is_aware(val):
                        val = _tz.localtime(val)
                    return val.strftime('%d/%m/%Y %H:%M')
                if isinstance(val, _dt.date):
                    return val.strftime('%d/%m/%Y')
            except Exception:
                return None
            return str(val)

        def _to_float(s):
            try:
                if s is None:
                    return None
                if isinstance(s, (int, float)):
                    return float(s)
                s = str(s).strip().replace(',', '.')
                return float(s) if s else None
            except Exception:
                return None

        def r3(v):
            try:
                return round(float(v), 3) if v is not None else None
            except Exception:
                return v

        # -------------------------
        # Iterate rows
        # -------------------------
        for row in qs.iterator(chunk_size=2000):
            dens = _to_float(row.get('densite_inspection'))
            temp = _to_float(row.get('temperature_inspection'))
            gov = _to_float(row.get('gov_total')) or 0
            gsv_val = _to_float(row.get('gsv_total')) or 0

            d15 = None
            vcf_val = None
            try:
                if dens is not None and temp is not None:
                    d15 = densite15(temp, dens)
                    if d15 is not None:
                        vcf_val = vcf(d15, temp)
                        d15 = round(float(d15), 5)
                    if vcf_val is not None:
                        vcf_val = round(float(vcf_val), 6)
                elif gov > 0 and gsv_val > 0:
                    vcf_val = round(gsv_val / gov, 6)
            except Exception:
                pass

            # ✅ "#.DOSSIER" must take value of numdos only
            db_numdos = row.get('numdos')
            dossier = '' if db_numdos is None else str(db_numdos)

            ws.append([
                _fmt_dt(row.get('dateheurecargaison')),
                row.get('nom_frontiere'),
                row.get('nom_importateur'),
                row.get('nom_entrepot'),
                row.get('nom_produit'),
                row.get('volume'),
                row.get('immatriculation'),
                row.get('declaration'),

                # ✅ only one dossier column
                dossier,

                _fmt_dt(row.get('requisitiondackdate')),
                _fmt_dt(row.get('date_echantillon')),
                _fmt_dt(row.get('date_reception_labo')),
                _fmt_dt(row.get('date_analyse')),
                _fmt_dt(row.get('date_inspection')),
                _fmt_dt(row.get('dateDechargement')),
                row.get('gov_total'),
                d15,
                temp,
                vcf_val,
                r3(row.get('mta_total')),
                r3(row.get('mtv_total')),
                r3(row.get('gsv_total')),
            ])

            done += 1
            if done % 1000 == 0 or done == total:
                try:
                    percent = round((done / max(total, 1)) * 100, 2)
                    self.update_state(
                        state='PROGRESS',
                        meta={'total': total, 'done': done, 'percent': percent}
                    )
                except Exception:
                    pass

        # -------------------------
        # Save to storage
        # -------------------------
        now_dt = timezone.now()
        now = now_dt.strftime('%Y%m%d%H%M%S')
        file_name = f"rapport_brut_{now}.xlsx"
        file_path = f"xlsx/{now_dt.strftime('%Y')}/{now_dt.strftime('%m')}/{file_name}"

        import io as _io
        content = _io.BytesIO()
        wb.save(content)
        content.seek(0)

        default_storage.save(file_path, ContentFile(content.read()))
        file_url = default_storage.url(file_path)

        return {'file_url': file_url, 'file_name': file_name}

    except Exception as e:
        # Let Celery mark the task as FAILURE by raising if you prefer,
        # but keeping your current pattern:
        return {'error': str(e)}
